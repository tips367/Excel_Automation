from openpyxl.chart import BarChart, Reference
import openpyxl as xl


def add_chart(filename):
    wb = xl.load_workbook(filename)
    sheet = wb['Sheet1']
    values_for_chart = Reference(sheet,
                          min_row=2,
                          max_row=sheet.max_row,
                          min_col=4,
                          max_col=4)
    chart = BarChart()
    chart.add_data(values_for_chart)
    sheet.add_chart(chart, 'F2')
    wb.save(filename)